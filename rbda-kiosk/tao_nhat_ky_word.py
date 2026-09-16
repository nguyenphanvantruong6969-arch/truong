#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Chuyển NHAT_KY_AI.html (bản artifact) thành tệp Word NHAT_KY_AI.docx.

Tệp Word giữ nguyên nội dung của artifact: từng câu lệnh (nguyên văn), mục đích
của câu lệnh đó, việc AI đã làm, tệp AI tạo ra và phần kiểm chứng. Thêm một bảng
tra nhanh ở đầu Mục 5 để dò theo số thứ tự câu lệnh.

Chạy lại:  python3 tao_nhat_ky_word.py [NHAT_KY_AI.html] [NHAT_KY_AI.docx]
"""

import re
import sys
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor

# ---------------------------------------------------------------- bảng màu
INK       = RGBColor(0x1A, 0x1D, 0x21)
INK_SOFT  = RGBColor(0x4A, 0x50, 0x59)
INK_FAINT = RGBColor(0x76, 0x7C, 0x85)
STAMP     = RGBColor(0x1F, 0x4E, 0x79)
FLAG      = RGBColor(0xA6, 0x3A, 0x2B)
CAUTION   = RGBColor(0x8A, 0x64, 0x10)
OK        = RGBColor(0x2F, 0x6B, 0x4F)

WASH_PROMPT  = "E7EEF5"
WASH_CAUTION = "F7EEDA"
WASH_FLAG    = "F7E7E4"
WASH_RULE    = "EAE7DE"
WASH_HEAD    = "F2F0EA"
WASH_OK      = "E6EFE9"

# Thanh của biểu đồ tiến độ vẽ bằng ký tự khối: 1 ô = 1 commit khi max <= 30.
BAR_CHAR, BAR_MAX_CHARS = "\u2588", 30

# Tông màu của hộp trong sơ đồ: (nền, màu viền, màu chữ nhấn)
TONE = {
    "is-stamp":   (WASH_PROMPT,  "1F4E79", STAMP),
    "is-ok":      (WASH_OK,      "2F6B4F", OK),
    "is-caution": (WASH_CAUTION, "8A6410", CAUTION),
    "is-flag":    (WASH_FLAG,    "A63A2B", FLAG),
    "":           (WASH_HEAD,    "DAD6CB", INK_SOFT),
}

F_BODY = "Calibri"
F_MONO = "Consolas"

PILL = {
    "allow": ("CHO PHÉP", OK),
    "cond":  ("CÓ ĐIỀU KIỆN", CAUTION),
    "deny":  ("KHÔNG ĐƯỢC PHÉP", FLAG),
}


# ------------------------------------------------------------ tiện ích XML
def _shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def shade_cell(cell, fill):
    _shade(cell._tc.get_or_add_tcPr(), fill)


def shade_par(par, fill):
    _shade(par._p.get_or_add_pPr(), fill)


def par_border(par, edges=("bottom",), size=6, color="DAD6CB", space=3):
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        bdr.append(el)


def cell_borders(cell, color="DAD6CB", size=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def cell_padding(cell, top=60, bottom=60, left=100, right=100):
    """Lề trong của một ô, đơn vị dxa (1/20 pt). Word mặc định 0 trên/dưới."""
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for canh, gt in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{canh}")
        el.set(qn("w:w"), str(gt))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def row_cant_split(row):
    """Không cắt một hộp làm đôi giữa hai trang."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def gop_o(cells, dau, rong):
    """Gộp `rong` ô liên tiếp từ cột `dau`, bỏ các đoạn rỗng thừa do gộp sinh ra."""
    cell = cells[dau]
    if rong > 1:
        cell = cell.merge(cells[min(dau + rong, len(cells)) - 1])
    pars = cell.paragraphs
    for par in pars[1:]:
        if not par.text.strip():
            par._p.getparent().remove(par._p)
    return cell


def css_var(tag, ten, mac_dinh=None):
    """Đọc một biến CSS (vd --gt) trong thuộc tính style= của một thẻ."""
    for phan in (tag.get("style") or "").split(";"):
        if ":" in phan:
            k, v = phan.split(":", 1)
            if k.strip() == ten:
                return float(v.strip())
    if mac_dinh is None:
        raise ValueError(f"thiếu biến {ten} trong style= của <{tag.name}>")
    return mac_dinh


def keep_with_next(par):
    par._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def add_page_field(par):
    """Chèn 'Trang X / Y' bằng trường của Word."""
    def field(instr):
        r = par.add_run()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), instr)
        r._r.addprevious(fld)

    run = par.add_run("Trang ")
    run.font.size = Pt(8)
    run.font.color.rgb = INK_FAINT
    field("PAGE")
    run = par.add_run(" / ")
    run.font.size = Pt(8)
    run.font.color.rgb = INK_FAINT
    field("NUMPAGES")


# ---------------------------------------------------------- chữ trong dòng
WS = re.compile(r"\s+")


def clean(text):
    return WS.sub(" ", text)


def style_run(run, bold=False, italic=False, mono=False, size=10.5,
              color=INK, underline=False):
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.size = Pt(size)
    run.font.color.rgb = color
    name = F_MONO if mono else F_BODY
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline(par, node, bold=False, italic=False, mono=False,
               size=10.5, color=INK):
    """Đổ nội dung trong dòng của một thẻ HTML vào đoạn văn Word."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = clean(str(child))
            if not text.strip() and not par.runs:
                continue
            if text:
                style_run(par.add_run(text), bold, italic, mono, size, color)
        elif isinstance(child, Tag):
            name = child.name
            if name == "br":
                par.add_run().add_break()
            elif name in ("strong", "b"):
                add_inline(par, child, True, italic, mono, size, color)
            elif name in ("em", "i"):
                add_inline(par, child, bold, True, mono, size, color)
            elif name == "code":
                add_inline(par, child, bold, italic, True, size - 0.5, color)
            elif name == "a":
                add_inline(par, child, bold, italic, mono, size, STAMP)
            else:
                add_inline(par, child, bold, italic, mono, size, color)


def add_text_par(doc_or_cell, node, size=10.5, color=INK, italic=False,
                 space_after=5, indent=0.0, align=None):
    par = doc_or_cell.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(0)
    if indent:
        par.paragraph_format.left_indent = Cm(indent)
    if align is not None:
        par.alignment = align
    add_inline(par, node, italic=italic, size=size, color=color)
    return par


def add_label(container, text, color=STAMP):
    par = container.add_paragraph()
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(2)
    style_run(par.add_run(text.upper()), bold=True, mono=True, size=8, color=color)
    keep_with_next(par)
    return par


def add_mono_block(container, text, fill=WASH_PROMPT, size=9.5, width=None):
    """Khối chữ máy (câu lệnh nguyên văn / số liệu thô) — bảng 1 ô có nền."""
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    if width is not None:
        table.autofit = False
        cell.width = width
    shade_cell(cell, fill)
    cell_borders(cell, color="DAD6CB")
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for i, line in enumerate(text.split("\n")):
        par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        style_run(par.add_run(line), mono=True, size=size, color=INK)
    return table


# ------------------------------------------------------------- khối lớn
def add_heading(doc, text, level):
    par = doc.add_paragraph()
    fmt = par.paragraph_format
    if level == 1:
        fmt.space_before = Pt(20)
        fmt.space_after = Pt(6)
        style_run(par.add_run(text), bold=True, size=16, color=INK)
        par_border(par, ("bottom",), size=12, color="1A1D21", space=4)
    elif level == 2:
        fmt.space_before = Pt(14)
        fmt.space_after = Pt(4)
        style_run(par.add_run(text), bold=True, size=12.5, color=STAMP)
    else:
        fmt.space_before = Pt(10)
        fmt.space_after = Pt(3)
        style_run(par.add_run(text), bold=True, size=11, color=INK)
    keep_with_next(par)
    return par


def add_bullets(container, ul, size=10, indent=0.6):
    for li in ul.find_all("li", recursive=False):
        par = container.add_paragraph()
        par.paragraph_format.left_indent = Cm(indent)
        par.paragraph_format.space_after = Pt(2)
        style_run(par.add_run("• "), size=size, color=STAMP)
        add_inline(par, li, size=size)


def render_alert(doc, div):
    caution = "is-caution" in (div.get("class") or [])
    fill = WASH_CAUTION if caution else WASH_FLAG
    color = CAUTION if caution else FLAG
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell_borders(cell, color="DAD6CB")
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for child in div.children:
        if not isinstance(child, Tag):
            continue
        if child.name == "h3":
            par = cell.add_paragraph()
            par.paragraph_format.space_after = Pt(4)
            add_inline(par, child, bold=True, size=11, color=color)
        elif child.name == "p":
            add_text_par(cell, child, size=10)
        elif child.name in ("span", "pre") and "quoted" in (child.get("class") or []):
            add_mono_block(cell, child.get_text().strip(), fill="FFFFFF", size=9)
        elif child.name == "ul":
            add_bullets(cell, child)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render_table(doc, table_tag):
    head = table_tag.find("thead")
    headers = [th for th in head.find_all("th")] if head else []
    body = table_tag.find("tbody")
    # Bảng không có <tbody> vẫn phải dựng được, đừng để sập cả lần chạy.
    body_rows = (body or table_tag).find_all("tr", recursive=False)
    if head is not None:
        body_rows = [r for r in body_rows if r.find_parent("thead") is None]
    ncols = len(headers) if headers else max(
        sum(int(td.get("colspan", 1)) for td in r.find_all(["td", "th"], recursive=False))
        for r in body_rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    if headers:
        cells = t.add_row().cells
        for cell, th in zip(cells, headers):
            shade_cell(cell, WASH_HEAD)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            add_inline(par, th, bold=True, size=9, color=INK_SOFT)
    for tr in body_rows:
        cells = t.add_row().cells
        cot = 0
        for td in tr.find_all(["td", "th"], recursive=False):
            rong = int(td.get("colspan", 1))
            if cot >= len(cells):
                break
            cell = gop_o(cells, cot, rong)
            if rong > 1:
                shade_cell(cell, WASH_HEAD)
            cot += rong
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            pill = td.find("span", class_="pill")
            if pill is not None:
                cls = [c for c in pill.get("class") if c in PILL]
                label, color = PILL.get(cls[0] if cls else "", ("", INK))
                style_run(par.add_run(label), bold=True, mono=True, size=8, color=color)
            else:
                add_inline(par, td, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# --------------------------------------------------- mục nhật ký câu lệnh
def parse_entry(article):
    gutter = article.find("div", class_="entry-gutter")
    seq = gutter.find("span", class_="seq").get_text(" ", strip=True)
    seq = clean(seq).replace(" ", "")
    when = clean(gutter.get_text(" ", strip=True).replace(
        gutter.find("span", class_="seq").get_text(" ", strip=True), "", 1)).strip()
    head = article.find("div", class_="entry-head")
    task = head.find("span", class_="task").get_text(" ", strip=True)
    pill = head.find("span", class_="pill")
    status = ("", INK)
    if pill is not None:
        cls = [c for c in pill.get("class") if c in PILL]
        status = PILL.get(cls[0] if cls else "", (pill.get_text(strip=True), INK))
    fields = []
    for field in article.find_all("div", class_="field"):
        label_tag = field.find("p", class_="field-label")
        label = label_tag.get_text(" ", strip=True) if label_tag else ""
        fields.append((label, field))
    return seq, when, clean(task), status, fields


def render_entry(doc, article):
    seq, when, task, (status_label, status_color), fields = parse_entry(article)

    head = doc.add_paragraph()
    head.paragraph_format.space_before = Pt(12)
    head.paragraph_format.space_after = Pt(3)
    style_run(head.add_run(f"#{seq}"), bold=True, mono=True, size=12, color=STAMP)
    style_run(head.add_run(f"  ·  {when}  ·  " if when else "  ·  "),
              mono=True, size=8.5, color=INK_FAINT)
    style_run(head.add_run(task), bold=True, size=11.5, color=INK)
    if status_label:
        style_run(head.add_run(f"   [{status_label}]"), bold=True, mono=True,
                  size=8, color=status_color)
    par_border(head, ("bottom",), size=6, color="1A1D21")
    keep_with_next(head)

    has_purpose = any(lbl.lower().startswith("mục đích") for lbl, _ in fields)
    if not has_purpose:
        add_label(doc, "Mục đích của câu lệnh")
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        style_run(par.add_run(task + "."), size=10.5, color=INK)

    for label, field in fields:
        if label:
            add_label(doc, label)
        for child in field.children:
            if not isinstance(child, Tag) or child is field.find("p", class_="field-label"):
                continue
            cls = child.get("class") or []
            if "field-label" in cls:
                continue
            if child.name == "pre":
                add_mono_block(doc, child.get_text().rstrip("\n"))
            elif child.name == "p":
                add_text_par(doc, child, size=10.5)
            elif child.name == "ul":
                mono = "files" in cls
                for li in child.find_all("li", recursive=False):
                    par = doc.add_paragraph()
                    par.paragraph_format.left_indent = Cm(0.6)
                    par.paragraph_format.space_after = Pt(1)
                    style_run(par.add_run("→ " if mono else "• "),
                              mono=mono, size=9.5, color=STAMP)
                    add_inline(par, li, mono=mono, size=9.5 if mono else 10.5)
            elif child.name == "table":
                render_table(doc, child)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ------------------------------------------------------- tờ chép tay Mục 8
def render_to_nk(doc, article):
    doc.add_page_break()
    head = article.find("div", class_="nk-dau")
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    cells_data = []
    for box in head.find_all("div", class_="nk-o"):
        nhan = box.find("span", class_="nk-nhan")
        label = nhan.get_text(" ", strip=True) if nhan else ""
        value = clean(box.get_text(" ", strip=True).replace(label, "", 1)).strip()
        wide = "nk-ngang" in (box.get("class") or [])
        cells_data.append((label, value, wide))

    def fill(cell, label, value):
        shade_cell(cell, WASH_RULE)
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        style_run(par.add_run(label), mono=True, size=7.5, color=INK_FAINT)
        par2 = cell.add_paragraph()
        par2.paragraph_format.space_after = Pt(0)
        style_run(par2.add_run(value if value else " " * 40), bold=True, size=10.5)

    i = 0
    while i < len(cells_data):
        label, value, wide = cells_data[i]
        row = t.add_row().cells
        if wide:
            merged = row[0].merge(row[1])
            fill(merged, label, value)
            i += 1
        else:
            fill(row[0], label, value)
            if i + 1 < len(cells_data) and not cells_data[i + 1][2]:
                fill(row[1], *cells_data[i + 1][:2])
                i += 2
            else:
                fill(row[1], "", "")
                i += 1

    for child in article.children:
        if not isinstance(child, Tag):
            continue
        cls = child.get("class") or []
        if "nk-dau" in cls:
            continue
        if child.name == "h4":
            so = child.find("span", class_="nk-so")
            num = so.get_text(strip=True) if so else ""
            text = clean(child.get_text(" ", strip=True))
            if num:
                text = text.replace(num, "", 1).strip()
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(9)
            par.paragraph_format.space_after = Pt(3)
            style_run(par.add_run(f"{num}. " if num else ""), bold=True,
                      mono=True, size=9.5, color=STAMP)
            style_run(par.add_run(text), bold=True, mono=True, size=9.5, color=STAMP)
            par_border(par, ("bottom",), size=4)
            keep_with_next(par)
        elif child.name == "ul":
            add_bullets(doc, child, size=10.5)
        elif child.name == "ol":
            add_bullets(doc, child, size=10.5)
        elif child.name == "pre":
            add_mono_block(doc, child.get_text().rstrip("\n"))
        elif child.name == "p":
            hoi = child.find("i", class_="nk-hoi")
            if hoi is not None:
                q = clean(hoi.get_text(" ", strip=True))
                hoi.extract()
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(1)
                style_run(par.add_run(q), mono=True, size=8.5, color=FLAG)
                add_text_par(doc, child, size=10.5)
            else:
                add_text_par(doc, child, size=10.5)
        elif "nk-ve" in cls:
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(6)
            par.paragraph_format.space_after = Pt(18)
            style_run(par.add_run(clean(child.get_text(" ", strip=True))),
                      italic=True, size=9.5, color=INK_FAINT)
            par_border(par, ("top", "bottom", "left", "right"), size=4)
        elif "nk-ky" in cls:
            sig = doc.add_table(rows=1, cols=2)
            for cell, box in zip(sig.rows[0].cells,
                                 child.find_all("div", class_="nk-ky-o")):
                p1 = cell.paragraphs[0]
                p1.paragraph_format.space_before = Pt(20)
                p1.paragraph_format.space_after = Pt(2)
                par_border(p1, ("bottom",), size=4, color="767C85")
                p2 = cell.add_paragraph()
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(p2.add_run(clean(box.get_text(" ", strip=True))),
                          size=9, color=INK_FAINT)
            doc.add_paragraph()


# ------------------------------------------------- sơ đồ và biểu đồ tiến độ
RONG_TRANG = Cm(17.0)          # bề ngang chữ với lề 2,0 cm hai bên


def tone_cua(tag):
    for lop in tag.get("class") or []:
        if lop in TONE:
            return TONE[lop]
    return TONE[""]


def kiem_tra_so_do(fig):
    """Sơ đồ sai cấu trúc thì dừng hẳn, đừng in ra một sơ đồ méo."""
    cot = int(fig.get("data-cot", 1))
    for lop in fig.find_all("div", class_="so-do-lop"):
        hop = lop.find_all("div", class_="so-do-hop", recursive=False)
        tong = sum(int(h.get("data-rong", 1)) for h in hop)
        if tong != cot:
            raise ValueError(
                f"sơ đồ '{lop.get('data-lop')}': tổng data-rong = {tong}, "
                f"khác data-cot = {cot}")
        for h in hop:
            if h.find(["pre", "table"]) is not None:
                raise ValueError(
                    f"sơ đồ '{lop.get('data-lop')}': hộp chứa <pre>/<table> "
                    "— sẽ sinh bảng lồng trong Word")
    return cot


def do_hop(cell, hop, nen, vien):
    """Đổ một .so-do-hop vào một ô Word: nền + viền theo tông, rồi ba đoạn chữ."""
    shade_cell(cell, nen)
    cell_borders(cell, color=vien)
    cell_padding(cell)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    phan = (("hop-ten", dict(bold=True, size=10, color=INK)),
            ("hop-tep", dict(mono=True, size=8.5, color=STAMP)),
            ("hop-mo", dict(size=8.5, color=INK_SOFT)))
    for lop, kieu in phan:
        the = hop.find(class_=lop)
        if the is None:
            continue
        par = cell.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(1)
        add_inline(par, the, **kieu)
    if not cell.paragraphs:
        cell.add_paragraph()


def render_so_do(doc, fig):
    """<figure class="so-do" data-cot="N"> → bảng Word 1 cột nhãn + N cột nội dung."""
    cot = kiem_tra_so_do(fig)
    cap = fig.find("figcaption")
    if cap is not None:
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.space_after = Pt(4)
        add_inline(par, cap, italic=True, size=9, color=INK_SOFT)
        keep_with_next(par)

    rong_nhan = Cm(3.0)
    rong_cot = Cm((17.0 - 3.0) / cot)
    t = doc.add_table(rows=0, cols=cot + 1)
    t.autofit = False

    for node in fig.children:
        if not isinstance(node, Tag):
            continue
        cls = node.get("class") or []
        if "so-do-lop" in cls:
            row = t.add_row()
            row_cant_split(row)
            cells = row.cells
            cells[0].width = rong_nhan
            for c in cells[1:]:
                c.width = rong_cot
            nhan = cells[0]
            shade_cell(nhan, WASH_RULE)
            cell_padding(nhan)
            par = nhan.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            style_run(par.add_run((node.get("data-lop") or "").upper()),
                      mono=True, size=7.5, color=INK_FAINT)
            vi_tri = 1
            for hop in node.find_all("div", class_="so-do-hop", recursive=False):
                rong = int(hop.get("data-rong", 1))
                nen, vien, _ = tone_cua(hop)
                o = gop_o(cells, vi_tri, rong)
                o.width = Cm(rong_cot.cm * rong)
                do_hop(o, hop, nen, vien)
                vi_tri += rong
        elif "so-do-mui-ten" in cls:
            row = t.add_row()
            row_cant_split(row)
            cells = row.cells
            cells[0].width = rong_nhan
            o = gop_o(cells, 1, cot)
            o.width = Cm(17.0 - 3.0)
            cell_padding(o, top=40, bottom=40)
            par = o.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            mau = OK if node.get("data-huong") == "len" else STAMP
            style_run(par.add_run(node.get_text(strip=True) or "\u2193"),
                      bold=True, mono=True, size=11, color=mau)
            style_run(par.add_run("   " + (node.get("data-nhan") or "")),
                      mono=True, size=8, color=INK_FAINT)

    ghi = fig.find("p", class_="so-do-ghi")
    if ghi is not None:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        add_inline(par, ghi, italic=True, size=9, color=INK_FAINT)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def kiem_tra_tien_do(fig):
    """Số vẽ ra thanh và số in bằng chữ phải khớp nhau, nếu không thì dừng."""
    hang = fig.find_all("li", class_="tien-do-hang")
    gia_tri = []
    for h in hang:
        gt = css_var(h, "--gt")
        vmax = css_var(h, "--max")
        the_so = h.find(class_="td-so")
        so_chu = float((the_so.get_text(strip=True) if the_so else "").replace(",", "."))
        ten = clean(h.find(class_="td-pha").get_text(" ", strip=True))
        if gt != so_chu:
            raise ValueError(f"biểu đồ '{ten}': --gt={gt} khác số in ra {so_chu}")
        gia_tri.append((gt, vmax, ten))
    dinh = max(g for g, _, _ in gia_tri)
    for gt, vmax, ten in gia_tri:
        if vmax != dinh:
            raise ValueError(f"biểu đồ '{ten}': --max={vmax} khác giá trị lớn nhất {dinh}")
    tong = sum(g for g, _, _ in gia_tri)
    khai = float(fig.get("data-tong-thay", tong))
    if tong != khai:
        raise ValueError(f"biểu đồ: cộng các giai đoạn = {tong}, khác data-tong-thay = {khai}")
    return hang, dinh, tong


def render_tien_do(doc, fig):
    """<figure class="tien-do"> → bảng 4 cột, thanh vẽ bằng ký tự khối."""
    hang, dinh, tong = kiem_tra_tien_do(fig)
    don_vi = fig.get("data-don-vi", "commit")
    moi_o = 1 if dinh <= BAR_MAX_CHARS else dinh / BAR_MAX_CHARS

    cap = fig.find("figcaption")
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(2)
    if cap is not None:
        add_inline(par, cap, italic=True, size=9, color=INK_SOFT)
    keep_with_next(par)
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    thang = (f"Mỗi ô {BAR_CHAR} = 1 {don_vi}." if moi_o == 1
             else f"Mỗi ô {BAR_CHAR} \u2248 {moi_o:.1f} {don_vi}.")
    style_run(par.add_run(thang + " Trục bắt đầu từ 0."), mono=True, size=8, color=INK_FAINT)
    keep_with_next(par)

    rong = [Cm(4.4), Cm(5.6), Cm(1.7), Cm(5.3)]
    t = doc.add_table(rows=0, cols=4)
    t.autofit = False
    tieu_de = t.add_row().cells
    for cell, chu, w in zip(tieu_de, ["GIAI ĐOẠN",
                                      f"PHÂN BỐ {don_vi.upper()} NHÌN THẤY ĐƯỢC (0–{dinh:g})",
                                      "SỐ", "VIỆC CHÍNH"], rong):
        cell.width = w
        shade_cell(cell, WASH_HEAD)
        cell_padding(cell)
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        style_run(par.add_run(chu), bold=True, mono=True, size=8, color=INK_SOFT)

    for h in hang:
        gt = css_var(h, "--gt")
        la_dinh = "la-dinh" in (h.get("class") or [])
        cells = t.add_row().cells
        for cell, w in zip(cells, rong):
            cell.width = w
            cell_padding(cell)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)

        # cột 1 — tên giai đoạn
        par = cells[0].paragraphs[0]
        stt = h.find(class_="td-stt")
        if stt is not None:
            style_run(par.add_run(stt.get_text(strip=True) + " · "),
                      bold=True, mono=True, size=9, color=STAMP)
            stt.extract()
        add_inline(par, h.find(class_="td-pha"), bold=True, size=9.5)

        # cột 2 — thanh + các ngày
        par = cells[1].paragraphs[0]
        n = int(round(gt / moi_o))
        if n > 0:
            style_run(par.add_run(BAR_CHAR * n), mono=True, size=8,
                      color=OK if la_dinh else STAMP)
        else:
            style_run(par.add_run("\u2014 "), mono=True, size=8, color=INK_FAINT)
            style_run(par.add_run(h.get("data-khong-do") or "không đo được"),
                      italic=True, size=8, color=CAUTION)
        ngay = h.find(class_="td-ngay")
        if ngay is not None:
            p2 = cells[1].add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            style_run(p2.add_run(clean(ngay.get_text(" ", strip=True))),
                      mono=True, size=7.5, color=INK_FAINT)

        # cột 3 — con số
        par = cells[2].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(par.add_run(f"{gt:g}"), bold=True, size=11, color=INK)
        p2 = cells[2].add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(0)
        style_run(p2.add_run(don_vi), size=7.5, color=INK_FAINT)

        # cột 4 — việc chính
        viec = h.find(class_="td-viec")
        if viec is not None:
            add_inline(cells[3].paragraphs[0], viec, size=9)

    meta = fig.get("data-tong-meta")
    cells = t.add_row().cells
    o = gop_o(cells, 0, 4)
    o.width = RONG_TRANG
    shade_cell(o, WASH_RULE)
    cell_padding(o)
    par = o.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    chu = f"Cộng các giai đoạn: {tong:g}"
    if meta:
        thieu = float(meta) - tong
        chu += (f"   ·   Tổng theo siêu dữ liệu: {float(meta):g}"
                f"   ·   Chưa nhìn thấy được: {thieu:g} (bản sao kho nông)")
    style_run(par.add_run(chu), mono=True, size=8, color=INK_SOFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ------------------------------------------------------------- bảng tra nhanh
def render_index(doc, entries):
    add_heading(doc, "Bảng tra nhanh — từng câu lệnh và mục đích", 2)
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(6)
    style_run(par.add_run(
        "Cột “Câu lệnh” chép rút gọn nguyên văn như đã gõ (kể cả lỗi chính tả); "
        "bản đầy đủ của từng câu lệnh nằm ngay bên dưới, theo đúng số thứ tự."),
        italic=True, size=9.5, color=INK_SOFT)

    t = doc.add_table(rows=0, cols=4)
    t.style = "Table Grid"
    widths = [Cm(1.6), Cm(2.2), Cm(6.6), Cm(6.6)]
    header = t.add_row().cells
    for cell, text, w in zip(header, ["STT", "Ngày", "Câu lệnh (rút gọn)",
                                      "Mục đích"], widths):
        shade_cell(cell, WASH_HEAD)
        cell.width = w
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(text), bold=True, mono=True, size=8, color=INK_SOFT)

    for article in entries:
        seq, when, task, (status_label, status_color), fields = parse_entry(article)
        prompt = ""
        for label, field in fields:
            if label.lower().startswith("câu lệnh"):
                pre = field.find("pre")
                if pre is not None:
                    prompt = clean(pre.get_text(" "))
                else:
                    # Mục không còn bản ghi nguyên văn — ghi đúng như vậy.
                    body = field.find("p", class_=lambda c: c != "field-label")
                    if body is None:
                        for cand in field.find_all("p"):
                            if "field-label" not in (cand.get("class") or []):
                                body = cand
                                break
                    if body is not None:
                        prompt = clean(body.get_text(" ", strip=True))
                break
        if len(prompt) > 150:
            prompt = prompt[:150].rstrip() + "…"
        purpose = task
        for label, field in fields:
            if label.lower().startswith("mục đích"):
                purpose = clean(field.get_text(" ", strip=True).replace(label, "", 1))
                break
        ngay = when.split("~")[0].strip() if when else ""
        cells = t.add_row().cells
        for cell, w in zip(cells, widths):
            cell.width = w
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        style_run(cells[0].paragraphs[0].add_run("#" + seq), bold=True,
                  mono=True, size=9, color=STAMP)
        style_run(cells[1].paragraphs[0].add_run(ngay), mono=True, size=8.5,
                  color=INK_SOFT)
        style_run(cells[2].paragraphs[0].add_run(prompt), mono=True, size=8.5)
        style_run(cells[3].paragraphs[0].add_run(purpose), size=9.5)
        if status_label:
            p = cells[3].add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(0)
            style_run(p.add_run(status_label), bold=True, mono=True, size=7.5,
                      color=status_color)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------- đầu trang
def render_masthead(doc, header):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    style_run(par.add_run(clean(header.find("p", class_="eyebrow").get_text()).upper()),
              bold=True, mono=True, size=9, color=STAMP)

    h1 = doc.add_paragraph()
    h1.paragraph_format.space_after = Pt(8)
    style_run(h1.add_run(clean(header.find("h1").get_text())), bold=True,
              size=22, color=INK)

    for p in header.find_all("p", class_="standfirst"):
        add_text_par(doc, p, size=10.5, color=INK_SOFT, space_after=6)

    meta = header.find("dl", class_="meta-grid")
    cells_data = [(clean(c.find("dt").get_text()), clean(c.find("dd").get_text()))
                  for c in meta.find_all("div", class_="meta-cell")]
    t = doc.add_table(rows=2, cols=len(cells_data))
    t.style = "Table Grid"
    for i, (dt, dd) in enumerate(cells_data):
        top = t.cell(0, i)
        shade_cell(top, WASH_HEAD)
        p = top.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(dt.upper()), mono=True, size=7.5, color=INK_FAINT)
        bottom = t.cell(1, i)
        p = bottom.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(dd), bold=True, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def normalize_cells(doc):
    """Word đòi mỗi ô bảng kết thúc bằng một đoạn văn, không phải bảng lồng."""
    def walk(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    children = list(cell._tc.iterchildren())
                    last = None
                    for child in children:
                        if child.tag.endswith("}tbl") or child.tag.endswith("}p"):
                            last = child
                    if last is not None and last.tag.endswith("}tbl"):
                        par = cell.add_paragraph()
                        par.paragraph_format.space_after = Pt(0)
                        par.paragraph_format.space_before = Pt(0)
                    walk(cell.tables)
    walk(doc.tables)
    body = doc.element.body
    last = None
    for child in body.iterchildren():
        if child.tag.endswith("}tbl") or child.tag.endswith("}p"):
            last = child
    if last is not None and last.tag.endswith("}tbl"):
        doc.add_paragraph()


# ----------------------------------------------------------------- chạy chính
def build(src: Path, dst: Path):
    soup = BeautifulSoup(src.read_text(encoding="utf-8"), "html.parser")
    sheet = soup.find("div", class_="sheet")

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = F_BODY
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), F_BODY)
    normal.paragraph_format.space_after = Pt(5)

    footer_par = section.footer.paragraphs[0]
    footer_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_par)

    entries = sheet.find_all("article", class_="entry")
    index_done = False
    pending_index = False
    bo_qua = []

    for node in sheet.children:
        if not isinstance(node, Tag):
            continue
        cls = node.get("class") or []
        name = node.name

        if "masthead" in cls:
            render_masthead(doc, node)
        elif name == "h2":
            num = node.find("span", class_="num")
            label = clean(num.get_text()) if num else ""
            text = clean(node.get_text(" ", strip=True))
            if label:
                text = text.replace(label, "", 1).strip()
                text = f"{label} — {text}"
            add_heading(doc, text, 1)
            # Neo theo id, không theo số mục: đánh số lại không phá được nữa.
            pending_index = node.get("id") == "muc-nhat-ky"
        elif name == "h3":
            add_heading(doc, clean(node.get_text(" ", strip=True)), 2)
        elif name == "h4":
            add_heading(doc, clean(node.get_text(" ", strip=True)), 3)
        elif name == "p" and "section-note" in cls:
            add_text_par(doc, node, size=10, color=INK_SOFT, space_after=8)
            if pending_index:
                render_index(doc, entries)
                add_heading(doc, "Nhật ký đầy đủ theo trình tự thời gian", 2)
                index_done, pending_index = True, False
        elif name == "p":
            add_text_par(doc, node)
        elif "alert" in cls:
            render_alert(doc, node)
        elif "table-wrap" in cls:
            render_table(doc, node.find("table"))
        elif "so-do" in cls:
            render_so_do(doc, node)
        elif "tien-do" in cls:
            render_tien_do(doc, node)
        elif name == "article" and "entry" in cls:
            if not index_done:
                render_index(doc, entries)
                add_heading(doc, "Nhật ký đầy đủ theo trình tự thời gian", 2)
                index_done = True
            render_entry(doc, node)
        elif name == "article" and "to-nk" in cls:
            render_to_nk(doc, node)
        elif name == "ol" and "todo" in cls:
            for i, li in enumerate(node.find_all("li", recursive=False), 1):
                box = li.find("span", class_="box")
                if box is not None:
                    box.extract()
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.8)
                par.paragraph_format.space_after = Pt(4)
                style_run(par.add_run(f"☐ {i}. "), bold=True, mono=True, size=9.5,
                          color=INK_SOFT)
                add_inline(par, li, size=10.5)
        elif name == "ul":
            add_bullets(doc, node)
        elif "signatures" in cls:
            doc.add_paragraph()
            t = doc.add_table(rows=1, cols=len(node.find_all("div", class_="sigline")))
            for cell, sig in zip(t.rows[0].cells,
                                 node.find_all("div", class_="sigline")):
                p1 = cell.paragraphs[0]
                p1.paragraph_format.space_before = Pt(24)
                par_border(p1, ("bottom",), size=4, color="767C85")
                cap = sig.find("span", class_="cap")
                p2 = cell.add_paragraph()
                style_run(p2.add_run(clean(cap.get_text()) if cap else ""),
                          mono=True, size=8, color=INK_FAINT)
                rest = clean(sig.get_text(" ", strip=True))
                if cap is not None:
                    rest = rest.replace(clean(cap.get_text()), "", 1).strip()
                if rest:
                    p3 = cell.add_paragraph()
                    style_run(p3.add_run(rest), size=9.5, color=INK_SOFT)
        elif name == "footer":
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(16)
            par_border(par, ("top",), size=6)
            add_inline(par, node, size=9, color=INK_FAINT)
        else:
            # Không nhận ra thì PHẢI kêu lên: bỏ qua im lặng là mất nguyên một
            # khối khỏi bản Word mà không ai biết.
            bo_qua.append(f"<{name} class={cls}>")

    normalize_cells(doc)

    props = doc.core_properties
    props.title = "Nhật ký sử dụng AI trong quá trình phát triển phần mềm"
    props.subject = "Phụ lục — Nhật ký câu lệnh AI · Đề tài Phân bổ Câu lạc bộ (RB-DA)"
    props.category = "Phụ lục nghiên cứu"
    props.comments = ("Sinh tự động từ NHAT_KY_AI.html bằng tao_nhat_ky_word.py — "
                      "nội dung giữ nguyên bản artifact.")

    doc.save(dst)
    if bo_qua:
        print("CẢNH BÁO — các khối cấp 1 bị bỏ qua, KHÔNG có trong bản Word:")
        for x in bo_qua:
            print("   ", x)
    return len(entries)


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else here / "NHAT_KY_AI.html"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else here / "NHAT_KY_AI.docx"
    n = build(src, dst)
    print(f"Đã ghi {dst} — {n} mục nhật ký.")
